import csv
import hashlib
import io
import json
import zipfile
from datetime import datetime
from decimal import Decimal
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from .db import Trade
from .schemas import TradeInput

ALIASES = {
 'exchange':['exchange','exch','exchange name'], 'segment':['segment','market segment'], 'expiry':['expiry','expiry date','expiry_date'],
 'lot_size':['lot_size','lot size'], 'strike':['strike','strike price'], 'option_type':['option_type','option type','ce/pe'],
 'symbol':['symbol','ticker','instrument','tradingsymbol'], 'entry_price':['entry_price','entry price','buy price','average buy price','open price'],
 'exit_price':['exit_price','exit price','sell price','average sell price','close price'], 'quantity':['quantity','qty','size','shares','contracts'],
 'entry_time':['entry_time','entry time','open time','date','entry date','buy date'], 'exit_time':['exit_time','exit time','close time','exit date','sell date'],
 'side':['side','direction','position'], 'commission':['commission','commissions','brokerage'], 'fees':['fees','fee','charges'],
 'asset_type':['asset_type','asset type','asset class'], 'status':['status'], 'setup':['setup','strategy'], 'notes':['notes','note'],
 'tags':['tags','tag'], 'multiplier':['multiplier','contract multiplier'], 'closed_quantity':['closed_quantity','closed quantity'],
 'emotion':['emotion','mood'], 'stop_loss':['stop_loss','stop loss','stop'], 'target_price':['target_price','target price','target'],
 'risk_amount':['risk_amount','risk amount','risk'], 'mark_price':['mark_price','mark price'], 'mfe':['mfe'], 'mae':['mae'], 'rating':['rating'],
 'planned_entry':['planned_entry','planned entry'], 'attributes':['attributes'],
}
NUMBERS={'entry_price','exit_price','quantity','commission','fees','multiplier','closed_quantity','stop_loss','target_price','risk_amount','mark_price','mfe','mae','rating','planned_entry','lot_size','strike'}


def read_table(raw, filename):
    if len(raw)>15*1024*1024:
        raise ValueError('File must be smaller than 15 MB')
    if filename.lower().endswith('.xlsx'):
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 60 * 1024 * 1024 or len(archive.infolist()) > 2000:
                raise ValueError('The expanded workbook exceeds the import limit')
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        iterator=wb.active.iter_rows(values_only=True)
        columns=[str(x or '').strip() for x in next(iterator)]
        rows=[]
        for line in iterator:
            if len(rows)>=20000:
                raise ValueError('Import at most 20,000 rows at a time')
            if any(x is not None for x in line):
                rows.append({k:(v.isoformat() if isinstance(v,datetime) else str(v) if v is not None else '') for k,v in zip(columns,line)})
        wb.close()
    else:
        text=raw.decode('utf-8-sig')
        try:
            dialect=csv.Sniffer().sniff(text[:4096],delimiters=',;\t')
        except csv.Error:
            dialect=csv.excel
        reader=csv.DictReader(io.StringIO(text),dialect=dialect)
        columns=reader.fieldnames or []
        rows=[]
        for row in reader:
            if len(rows)>=20000:
                raise ValueError('Import at most 20,000 rows at a time')
            rows.append(row)
    if len(columns) > 100 or any(len(str(cell)) > 20000 for row in rows for cell in row.values()):
        raise ValueError('The file has too many columns or oversized cells')
    if not columns or not rows:
        raise ValueError('The file has no data rows')
    mapping={}
    for field,names in ALIASES.items():
        found=next((c for c in columns if c.lower().strip() in names),None)
        if found:
            mapping[field]=found
    return columns,rows,mapping


def fingerprint(data):
    # Trade identity ignores notes, marks, and strategy annotations.
    keys=['account_id','symbol','asset_type','exchange','segment','expiry','strike','option_type','lot_size','side','entry_time','exit_time','entry_price','exit_price','quantity','closed_quantity','multiplier','commission','fees']
    # Database reads may return floats (0.0), whereas validated defaults
    # can be ints (0). Normalize numerics so a saved row retains its identity.
    canonical={k:(format(Decimal(str(data[k])).normalize(),'f') if k in NUMBERS and data.get(k) is not None else data.get(k)) for k in keys}
    return hashlib.sha256(json.dumps(canonical,sort_keys=True).encode()).hexdigest()


def preview_rows(db, rows, mapping, account_id, defaults=None):
    defaults=defaults or {}
    seen={fingerprint({c.name:getattr(t,c.name) for c in t.__table__.columns}) for t in db.scalars(select(Trade).where(Trade.account_id==account_id))}
    result=[]
    for i,row in enumerate(rows):
        try:
            data={**defaults,'account_id':account_id}
            for key,column in mapping.items():
                val=row.get(column)
                if val is None or str(val).strip()=='':
                    continue
                val=str(val).strip()
                if key in NUMBERS:
                    val=float(val.replace(',','').replace('$','').replace('₹',''))
                elif key=='tags':
                    val=json.loads(val) if val.startswith('[') else [t.strip() for t in val.replace(';',',').split(',') if t.strip()]
                elif key=='attributes':
                    val=json.loads(val)
                if key in ('entry_time','exit_time','expiry'):
                    val=normalize_date(val)
                data[key]=val
            data['side']={'buy':'Long','sell':'Short','long':'Long','short':'Short'}.get(str(data.get('side','Long')).lower(),data.get('side'))
            for key in ['status','asset_type','emotion']:
                if key in data:
                    data[key]=str(data[key]).title()
            if data.get('option_type'):
                data['option_type']={'CALL':'CE','PUT':'PE','C':'CE','P':'PE'}.get(str(data['option_type']).upper(),str(data['option_type']).upper())
            if not data.get('status'):
                data['status']='Closed' if data.get('exit_price') is not None else 'Open'
            if 'multiplier' not in data and data.get('asset_type') in ('Futures','Options'):
                raise ValueError('Provide a contract multiplier for options/futures; it cannot be inferred safely')
            normalized=TradeInput.model_validate(data).model_dump()
            fp=fingerprint(normalized)
            duplicate=fp in seen
            seen.add(fp)
            result.append({'row':i+2,'status':'duplicate' if duplicate else 'valid','data':normalized,'fingerprint':fp})
        except (ValidationError,ValueError,TypeError) as e:
            msg='; '.join(f"{'.'.join(str(v) for v in err['loc'])}: {err['msg']}" for err in e.errors()) if isinstance(e,ValidationError) else str(e)
            result.append({'row':i+2,'status':'error','error':msg})
    return result


def normalize_date(value):
    value=str(value).strip()
    try:
        datetime.fromisoformat(value.replace('Z','+00:00'))
        return value
    except ValueError:
        for fmt in ('%d/%m/%Y %H:%M:%S','%d/%m/%Y %H:%M','%d-%m-%Y %H:%M:%S','%d-%m-%Y %H:%M','%d/%m/%Y','%d-%m-%Y','%d-%b-%Y'):
            try:
                return datetime.strptime(value,fmt).isoformat()
            except ValueError:
                pass
    raise ValueError('Use ISO dates or Indian DD/MM/YYYY HH:mm:ss dates')
